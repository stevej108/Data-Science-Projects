import re
import json
import hashlib
from docx.oxml.shared import OxmlElement, qn
from docx.opc.constants import RELATIONSHIP_TYPE
from pathlib import Path

# -------------------------
# TEXT CLEANING
# -------------------------

def normalize_text(text: str) -> str:
    if not text:
        return ""

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


# -------------------------
# CHUNK BUILDER
# -------------------------

def build_chunk(document: dict, chunk_number: int, text: str,
                section_title: str = "", metadata: dict | None = None):

    metadata = metadata or {}

    text = normalize_text(text)
    section_title = normalize_text(section_title)

    combined_text = "\n".join(x for x in [section_title, text] if x)

    chunk_id = hashlib.md5(
        f"{document['document_id']}_{chunk_number}".encode()
    ).hexdigest()

    return {
        "chunk_id": chunk_id,
        "document_id": document["document_id"],
        "document_type": document["document_type"],
        "file_name": document["file_name"],
        "path": document["path"],
        "relative_path": document["relative_path"],
        "extension": document["extension"],
        "chunk_number": chunk_number,
        "section_title": section_title,
        "text": text,
        "combined_text": combined_text,
        "metadata": metadata,
        "modified": document["modified"],
        "created": document["created"],
        "size_mb": document["size_mb"],
        "onedrive_url": document["onedrive_url"],
        "source": {
            "document": document["file_name"],
            "relative_path": document["relative_path"],
            "chunk_number": chunk_number,
        },
        "retrieval_score": None,
        "rerank_score": None,
    }


# -------------------------
# JSON extraction (OK but tighten regex)
# -------------------------

def extract_json(text: str):
    text = text.strip().replace("```json", "").replace("```", "")

    try:
        return json.loads(text)
    except Exception:
        pass

    match = re.search(r"\{[\s\S]*\}", text)
    if match:
        return json.loads(match.group())

    raise ValueError("No JSON object found")


# ==========================================================
# PowerPoint Helper - Extract Slide Text
# ==========================================================

def extract_ppt_text(slide):
    """
    Extract all searchable text from a PowerPoint slide.
    """

    text = []

    for shape in slide.shapes:

        # Text boxes, placeholders, etc.
        if hasattr(shape, "text"):

            value = normalize_text(shape.text)

            if value:
                text.append(value)

        # Tables
        elif getattr(shape, "has_table", False):

            for row in shape.table.rows:

                cells = [
                    normalize_text(cell.text)
                    for cell in row.cells
                    if normalize_text(cell.text)
                ]

                if cells:
                    text.append(" | ".join(cells))

    return "\n".join(text)


# ==========================================================
# PowerPoint Helper - Extract Speaker Notes
# ==========================================================

def extract_ppt_notes(slide):
    """
    Extract speaker notes if present.
    """

    try:

        notes = []

        for shape in slide.notes_slide.shapes:

            if hasattr(shape, "text"):

                value = normalize_text(shape.text)

                if value:
                    notes.append(value)

        return "\n".join(notes)

    except Exception:

        return ""